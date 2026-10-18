"""Text and Markdown to PDF / DOCX."""

from __future__ import annotations

import html
from pathlib import Path

import pymupdf

from ..utils import output_path

_CSS = """
body { font-family: sans-serif; font-size: 11pt; }
h1 { font-size: 20pt; } h2 { font-size: 16pt; } h3 { font-size: 13pt; }
code, pre { font-family: monospace; background-color: #f0f0f0; }
table { border-collapse: collapse; } td, th { border: 1px solid #999; padding: 3px; }
"""


def read_text(src: Path) -> str:
    raw = src.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "cp1254", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def html_to_pdf(body: str, dest: Path) -> Path:
    """Lay out an HTML fragment on A4 pages with PyMuPDF's Story engine."""
    story = pymupdf.Story(html=body, user_css=_CSS)
    writer = pymupdf.DocumentWriter(str(dest))
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (50, 50, -50, -50)
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()
    return dest


def text_to_pdf(src: Path) -> list[Path]:
    text = read_text(src)
    if src.suffix.lower() in (".md", ".markdown"):
        import markdown

        body = markdown.markdown(text, extensions=["tables", "fenced_code"])
    else:
        body = f"<pre style='font-family: monospace; font-size: 10pt'>{html.escape(text)}</pre>"
    return [html_to_pdf(body, output_path(src, ".pdf"))]


def text_to_docx(src: Path) -> list[Path]:
    from docx import Document

    doc = Document()
    text = read_text(src)
    is_md = src.suffix.lower() in (".md", ".markdown")
    for line in text.splitlines():
        stripped = line.lstrip("#").strip()
        if is_md and line.startswith("#"):
            doc.add_heading(stripped, level=min(len(line) - len(line.lstrip("#")), 4))
        else:
            doc.add_paragraph(line)
    dest = output_path(src, ".docx")
    doc.save(dest)
    return [dest]

"""DOCX conversions."""

from __future__ import annotations

import html
from pathlib import Path

from ..utils import output_path


def docx_to_txt(src: Path) -> list[Path]:
    from docx import Document

    doc = Document(str(src))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    dest = output_path(src, ".txt")
    dest.write_text("\n".join(lines), encoding="utf-8")
    return [dest]


def docx_to_pdf(src: Path) -> list[Path]:
    """Render paragraphs, headings, bold/italic runs, lists and tables to PDF (layout is simplified)."""
    from docx import Document

    from .text import html_to_pdf

    doc = Document(str(src))
    parts: list[str] = []

    def runs_html(paragraph) -> str:
        out = []
        for run in paragraph.runs:
            text = html.escape(run.text)
            if run.bold:
                text = f"<b>{text}</b>"
            if run.italic:
                text = f"<i>{text}</i>"
            if run.underline:
                text = f"<u>{text}</u>"
            out.append(text)
        return "".join(out)

    for paragraph in doc.paragraphs:
        content = runs_html(paragraph)
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading") or style == "title":
            level = style.replace("heading", "").strip()
            tag = f"h{level}" if level.isdigit() and 1 <= int(level) <= 3 else "h1" if style == "title" else "h3"
            parts.append(f"<{tag}>{content}</{tag}>")
        elif "list" in style:
            parts.append(f"<p>&bull; {content}</p>")
        else:
            parts.append(f"<p>{content or '&nbsp;'}</p>")
    for table in doc.tables:
        rows = "".join(
            "<tr>" + "".join(f"<td>{html.escape(cell.text)}</td>" for cell in row.cells) + "</tr>" for row in table.rows
        )
        parts.append(f"<table>{rows}</table>")

    dest = output_path(src, ".pdf")
    html_to_pdf("".join(parts) or "<p></p>", dest)
    return [dest]
